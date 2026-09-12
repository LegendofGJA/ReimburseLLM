"""
Aplikasi Scan & Ekstraksi Reimburse Struk
==========================================
Upload foto struk (bensin, parkir, belanja Teazzi/drink) sekaligus untuk satu
periode, aplikasi akan:
  1. Auto-ping provider AI vision (Kagiro / Bandel) dan pilih model yang aktif.
  2. Ekstrak data dari tiap struk (tanggal, kategori, deskripsi, nominal)
     mengikuti aturan bisnis per kategori.
  3. Urutkan hasil secara kronologis berdasarkan TANGGAL STRUK (bukan urutan upload).
  4. Isi ke template Excel FORM_REIMBURSE (baris 12 dst, total otomatis di E62/ dst).
  5. Gabungkan seluruh foto struk asli menjadi satu PDF TANPA kompresi/downsizing.

Jalankan dengan:
    pip install -r requirements.txt
    streamlit run app.py
"""

import base64
import io
import json
import re
from datetime import date, datetime
from calendar import monthrange

import img2pdf
import openpyxl
import pandas as pd
import requests
import streamlit as st
from dateutil import parser as dateparser
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Scan Reimburse Struk", page_icon="🧾", layout="wide")

st.markdown(
    """
    <h2>🧾 Aplikasi Scan & Ekstraksi Reimburse Struk</h2>
    <p>Upload foto struk bensin, parkir, dan Teazzi (drink) sekaligus untuk satu bulan.
    Sistem akan mengekstrak data, mengurutkannya berdasarkan tanggal transaksi,
    mengisinya ke template Excel, dan menggabungkan foto struk asli ke satu PDF
    tanpa kompresi.</p>
    """,
    unsafe_allow_html=True,
)

# ─────────────────────────────────────────────────────────────────────────
# KONFIGURASI PROVIDER
# API key TIDAK ditulis langsung di kode (aman untuk repo publik/privat).
# Diambil dari st.secrets (file .streamlit/secrets.toml lokal, atau menu
# "Secrets" di Streamlit Community Cloud) dengan fallback ke environment
# variable. Lihat .streamlit/secrets.toml.example untuk formatnya.
# ─────────────────────────────────────────────────────────────────────────
import os


def _get_secret(key: str, default: str = "") -> str:
    try:
        if key in st.secrets:
            return st.secrets[key]
    except Exception:
        pass  # secrets.toml belum ada / tidak terbaca -> lanjut ke env var
    return os.environ.get(key, default)


PROVIDERS = {
    "Kagiro": {
        "base_url": _get_secret("KAGIRO_BASE_URL", "https://api.kagiro.net/v1"),
        "api_key": _get_secret("KAGIRO_API_KEY"),
    },
    "Bandel": {
        "base_url": _get_secret("BANDEL_BASE_URL", "https://bandelbanget.xyz/v1"),
        "api_key": _get_secret("BANDEL_API_KEY"),
    },
}
# Provider tanpa API key terisi otomatis di-skip saat auto-ping, supaya tidak
# muncul error membingungkan kalau salah satu provider memang belum dikonfigurasi.
PROVIDERS = {name: cfg for name, cfg in PROVIDERS.items() if cfg["api_key"]}

if not PROVIDERS:
    st.error(
        "Belum ada API key yang dikonfigurasi. Buat file `.streamlit/secrets.toml` "
        "(lihat `.streamlit/secrets.toml.example`) atau set environment variable "
        "KAGIRO_API_KEY / BANDEL_API_KEY, lalu jalankan ulang aplikasi."
    )
    st.stop()

TEMPLATE_PATH = "FORM_REIMBURSE_template.xlsx"
DATA_START_ROW = 12          # baris pertama tabel data di template
DATA_END_ROW_DEFAULT = 61    # baris terakhir yang sudah disiapkan template
TOTAL_ROW_DEFAULT = 62       # baris 'Total' di template


def test_ping_and_get_models(provider_name: str):
    """Ping provider dan ambil daftar model vision yang tersedia."""
    cfg = PROVIDERS[provider_name]
    url = f"{cfg['base_url']}/models"
    headers = {"Authorization": f"Bearer {cfg['api_key']}"}
    try:
        response = requests.get(url, headers=headers, timeout=8)
        if response.status_code == 200:
            data = response.json()
            models = data.get("data", data)
            model_list = [m.get("id") for m in models if isinstance(m, dict) and "id" in m]
            if not model_list:
                model_list = ["qwen-vl-max", "gpt-4o"]
            return "✅ Berhasil terhubung (HTTP 200)", model_list
        return f"⚠️ Respon gagal: Status {response.status_code}", []
    except Exception as e:
        return f"❌ Gagal terhubung / offline ({e})", []


def auto_select_provider():
    """Coba tiap provider berurutan, pakai yang pertama berhasil (auto-fallback)."""
    results = {}
    for name in PROVIDERS:
        status, models = test_ping_and_get_models(name)
        results[name] = (status, models)
        if models:
            return name, status, models, results
    # semua gagal -> tetap kembalikan provider pertama untuk ditampilkan statusnya
    first = list(PROVIDERS.keys())[0]
    status, models = results[first]
    return first, status, models, results


def encode_image(file_bytes: bytes) -> str:
    return base64.b64encode(file_bytes).decode("utf-8")


EXTRACTION_PROMPT = """Kamu adalah sistem OCR untuk struk belanja Indonesia.
Analisis gambar struk ini dan kembalikan HANYA JSON murni (tanpa markdown, tanpa teks tambahan) dengan struktur persis berikut:

{
  "date": "YYYY-MM-DD",
  "type": "bensin" | "parkir" | "drink",
  "nominal": <angka total akhir struk, tanpa titik/koma/Rp>,
  "fuel_type": "<nama jenis BBM apa adanya di struk, contoh: Pertalite / Pertamax / Pertamax Turbo, atau null jika bukan bensin>",
  "liters": <jumlah liter sebagai angka, atau null jika tidak ada / bukan bensin>,
  "drink_name": "<jenis/menu minuman yang dibeli, atau null jika bukan drink>",
  "outlet_name": "<nama toko/outlet, contoh Teazzi, atau null jika bukan drink>",
  "location_name": "<nama tempat/lokasi parkir, atau null jika bukan parkir. Jika parkir tapi nama tempat tidak tertera, isi dengan '-'>"
}

Aturan klasifikasi "type":
- Struk pom bensin / SPBU -> "bensin"
- Struk parkir -> "parkir"
- Struk Teazzi atau minuman lain -> "drink"

Ambil "nominal" sebagai TOTAL AKHIR yang benar-benar dibayar pada struk.
Tanggal WAJIB diambil dari tanggal transaksi yang tertera di struk, bukan diasumsikan.
Kembalikan JSON murni saja, tidak ada teks lain sebelum atau sesudahnya."""


def call_vision_api(provider_name: str, model: str, image_bytes: bytes) -> dict:
    cfg = PROVIDERS[provider_name]
    api_url = f"{cfg['base_url']}/chat/completions"
    headers = {
        "Authorization": f"Bearer {cfg['api_key']}",
        "Content-Type": "application/json",
    }
    base64_img = encode_image(image_bytes)
    payload = {
        "model": model,
        "messages": [
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": EXTRACTION_PROMPT},
                    {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{base64_img}"}},
                ],
            }
        ],
        "temperature": 0.1,
    }
    res = requests.post(api_url, headers=headers, json=payload, timeout=45)
    res.raise_for_status()
    content_text = res.json()["choices"][0]["message"]["content"]
    content_text = re.sub(r"```json|```", "", content_text).strip()
    return json.loads(content_text)


def build_description(item: dict) -> str:
    """Susun kolom description sesuai aturan bisnis per kategori (dipastikan di
    Python, tidak hanya mengandalkan kepatuhan model AI terhadap instruksi)."""
    itype = (item.get("type") or "").strip().lower()

    if itype == "bensin":
        fuel = (item.get("fuel_type") or "").strip()
        liters = item.get("liters")
        if not fuel:
            return "-"
        if "pertalite" in fuel.lower():
            return fuel
        if liters:
            try:
                liters_fmt = f"{float(liters):g}"
            except (TypeError, ValueError):
                liters_fmt = str(liters)
            return f"{fuel} {liters_fmt} Liter"
        return fuel

    if itype == "drink":
        drink = (item.get("drink_name") or "").strip()
        outlet = (item.get("outlet_name") or "").strip()
        parts = [p for p in [drink, outlet] if p]
        return " - ".join(parts) if parts else "-"

    if itype == "parkir":
        loc = (item.get("location_name") or "").strip()
        return loc if loc else "-"

    return "-"


def parse_date_safe(raw):
    if not raw:
        return None
    try:
        return dateparser.parse(str(raw), dayfirst=False, yearfirst=True).date()
    except Exception:
        try:
            return dateparser.parse(str(raw), dayfirst=True).date()
        except Exception:
            return None


def build_rows(extracted_items: list) -> pd.DataFrame:
    rows = []
    for item in extracted_items:
        d = parse_date_safe(item.get("date"))
        try:
            nominal = float(re.sub(r"[^\d.\-]", "", str(item.get("nominal", 0))) or 0)
        except ValueError:
            nominal = 0.0
        rows.append(
            {
                "date": d,
                "category": item.get("type", "-"),
                "description": build_description(item),
                "nominal": nominal,
            }
        )
    df = pd.DataFrame(rows)
    # struk tanpa tanggal terbaca ditaruh paling akhir, bukan hilang
    df["_sort_key"] = df["date"].apply(lambda d: d if d else date.max)
    df = df.sort_values(by="_sort_key", ascending=True).drop(columns=["_sort_key"]).reset_index(drop=True)
    return df


def fill_excel_template(df: pd.DataFrame, template_path: str, header_info: dict) -> bytes:
    wb = openpyxl.load_workbook(template_path)
    ws = wb["FORM"]

    # Info header (opsional override)
    if header_info.get("name"):
        ws["C6"] = header_info["name"]
    if header_info.get("department"):
        ws["C7"] = header_info["department"]
    if header_info.get("purpose"):
        ws["C8"] = header_info["purpose"]
    if header_info.get("bank_acc"):
        ws["C9"] = header_info["bank_acc"]

    # Periode awal - akhir bulan di E6 & E7
    valid_dates = [d for d in df["date"] if d]
    if valid_dates:
        ref = min(valid_dates)
    else:
        ref = date.today()
    period_start = date(ref.year, ref.month, 1)
    period_end = date(ref.year, ref.month, monthrange(ref.year, ref.month)[1])
    ws["E6"].number_format = "d mmm yyyy"
    ws["E7"].number_format = "d mmm yyyy"
    ws["E6"] = period_start
    ws["E7"] = period_end

    n = len(df)
    available_rows = DATA_END_ROW_DEFAULT - DATA_START_ROW + 1  # 50 baris tersedia di template

    if n > available_rows:
        extra = n - available_rows
        ws.insert_rows(DATA_END_ROW_DEFAULT + 1, amount=extra)
        # salin format dari baris terakhir template ke baris baru
        for i in range(extra):
            src_row = DATA_END_ROW_DEFAULT
            dst_row = DATA_END_ROW_DEFAULT + 1 + i
            for col in range(2, 6):  # B..E
                src_cell = ws.cell(row=src_row, column=col)
                dst_cell = ws.cell(row=dst_row, column=col)
                dst_cell.number_format = src_cell.number_format
                dst_cell.font = src_cell.font.copy()
                dst_cell.border = src_cell.border.copy()
                dst_cell.fill = src_cell.fill.copy()
                dst_cell.alignment = src_cell.alignment.copy()
        total_row = TOTAL_ROW_DEFAULT + extra
        data_end_row = DATA_END_ROW_DEFAULT + extra
    else:
        total_row = TOTAL_ROW_DEFAULT
        data_end_row = DATA_END_ROW_DEFAULT

    for i, row in df.iterrows():
        r = DATA_START_ROW + i
        ws.cell(row=r, column=2, value=row["date"])          # B: Date
        ws.cell(row=r, column=3, value=row["category"])       # C: Category
        ws.cell(row=r, column=4, value=row["description"])    # D: Description
        ws.cell(row=r, column=5, value=row["nominal"])         # E: Amount

    total_cell = ws.cell(row=total_row, column=5)
    total_cell.value = f"=SUM(E{DATA_START_ROW}:E{data_end_row})"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def merge_images_to_pdf(image_bytes_list: list) -> bytes:
    """Gabungkan gambar asli ke satu PDF tanpa kompresi/downsizing (img2pdf
    menyisipkan data gambar apa adanya, tidak melakukan re-encode)."""
    return img2pdf.convert(image_bytes_list)


# ─────────────────────────────────────────────────────────────────────────
# UI
# ─────────────────────────────────────────────────────────────────────────
with st.sidebar:
    st.subheader("Data Pemohon")
    name = st.text_input("Name", value="Gabriel Jorma Abraham")
    department = st.text_input("Department", value="QC")
    purpose = st.text_input("Purpose", value="Reimburse")
    bank_acc = st.text_input("Bank Acc.", value="BCA: Gabriel Jorma Abraham 2770715752")
    st.caption("Periode (kolom E6 & E7) diisi otomatis: awal & akhir bulan dari tanggal struk yang ter-upload.")

col1, col2 = st.columns([1, 1])
with col1:
    mode = st.radio("Pemilihan provider OCR", ["Auto (coba semua, pakai yang aktif)", "Pilih manual"], horizontal=False)

if mode.startswith("Auto"):
    provider_name, status_ping, available_models, all_status = auto_select_provider()
    st.markdown(f"**Provider aktif:** {provider_name} — {status_ping}")
    with st.expander("Status semua provider"):
        for pname, (pstatus, _) in all_status.items():
            st.write(f"- {pname}: {pstatus}")
else:
    provider_name = st.selectbox("Pilih Provider OCR", list(PROVIDERS.keys()))
    status_ping, available_models = test_ping_and_get_models(provider_name)
    st.markdown(f"**Status Ping Provider:** {status_ping}")

selected_model = st.selectbox("Pilih Model Vision", available_models if available_models else ["Model tidak tersedia"])

uploaded_files = st.file_uploader(
    "Unggah foto-foto struk (Bensin, Parkir, Teazzi) sekaligus untuk sebulan",
    type=["png", "jpg", "jpeg"],
    accept_multiple_files=True,
)

st.info(
    """
**Aturan ekstraksi & format output:**
1. **Sorting kronologis:** hasil di Excel diurutkan dari tanggal transaksi struk paling awal ke akhir (bukan urutan upload).
2. **Bensin:** Pertalite → description cukup nama BBM. Pertamax/jenis lain → description berisi jenis BBM + jumlah liter, nominal tetap total akhir struk.
3. **Drink (Teazzi):** description berisi jenis minuman & nama outlet, nominal total akhir.
4. **Parkir:** description berisi nama tempat (atau `-` jika tidak ada), nominal total.
5. **Periode E6/E7:** otomatis diisi awal & akhir bulan berdasarkan tanggal struk.
6. **PDF gabungan:** seluruh foto struk digabung ke satu PDF **tanpa kompresi/downsizing**.
    """
)

if "extracted_items" not in st.session_state:
    st.session_state.extracted_items = []
if "image_bytes_list" not in st.session_state:
    st.session_state.image_bytes_list = []

if st.button("🚀 Mulai Proses OCR, Sorting, Generate Excel & PDF", type="primary"):
    if not uploaded_files:
        st.warning("Silakan unggah minimal satu foto struk terlebih dahulu.")
    elif not available_models or selected_model == "Model tidak tersedia":
        st.error("Provider tidak aktif atau model tidak ditemukan. Gagal memproses.")
    else:
        extracted_items = []
        image_bytes_list = []
        failures = []

        progress_bar = st.progress(0)
        total_files = len(uploaded_files)

        with st.spinner("Sedang memproses ekstraksi gambar dengan AI Vision..."):
            for i, file in enumerate(uploaded_files):
                img_bytes = file.getvalue()
                image_bytes_list.append(img_bytes)
                try:
                    parsed = call_vision_api(provider_name, selected_model, img_bytes)
                    extracted_items.append(parsed)
                except Exception as e:
                    failures.append((file.name, str(e)))
                progress_bar.progress((i + 1) / total_files)

        if failures:
            with st.expander(f"⚠️ {len(failures)} file gagal diproses"):
                for fname, err in failures:
                    st.write(f"- {fname}: {err}")

        if extracted_items:
            st.session_state.extracted_items = extracted_items
            st.session_state.image_bytes_list = image_bytes_list
            st.success(f"Ekstraksi selesai: {len(extracted_items)} struk berhasil diproses.")
        else:
            st.error("Tidak ada struk yang berhasil diekstrak.")

if st.session_state.extracted_items:
    df = build_rows(st.session_state.extracted_items)
    display_df = df.copy()
    display_df["date"] = display_df["date"].apply(lambda d: d.strftime("%d %b %Y") if d else "-")
    st.dataframe(display_df, use_container_width=True)
    st.markdown(f"**Total: Rp {df['nominal'].sum():,.0f}**".replace(",", "."))

    excel_bytes = fill_excel_template(
        df,
        TEMPLATE_PATH,
        {"name": name, "department": department, "purpose": purpose, "bank_acc": bank_acc},
    )
    pdf_bytes = merge_images_to_pdf(st.session_state.image_bytes_list)

    dcol1, dcol2 = st.columns(2)
    with dcol1:
        st.download_button(
            "📥 Download Excel Reimburse",
            data=excel_bytes,
            file_name=f"FORM_REIMBURSE_{date.today().strftime('%Y%m')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    with dcol2:
        st.download_button(
            "📥 Download PDF Struk Gabungan (tanpa kompresi)",
            data=pdf_bytes,
            file_name=f"Struk_Gabungan_{date.today().strftime('%Y%m')}.pdf",
            mime="application/pdf",
        )
