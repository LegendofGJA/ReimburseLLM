"""
excel_core.py — Isi data ke template Excel FORM_REIMBURSE.

Menangani:
  - header info (Name/Department/Purpose/Bank Acc),
  - periode E6/E7 (awal & akhir bulan dari tanggal struk),
  - penulisan baris data mulai DATA_START_ROW,
  - perluasan baris otomatis bila data > baris tersedia,
  - total otomatis (formula SUM).
"""

import io
from calendar import monthrange
from datetime import date

import openpyxl

TEMPLATE_PATH = "FORM_REIMBURSE_template.xlsx"
DATA_START_ROW = 12
DATA_END_ROW_DEFAULT = 61
TOTAL_ROW_DEFAULT = 62
SHEET_NAME = "FORM"


def _guess_sheet(ws):
    return ws


def fill_excel_template(df, template_path: str, header_info: dict) -> bytes:
    wb = openpyxl.load_workbook(template_path)
    ws = wb[SHEET_NAME]

    if header_info.get("name"):
        ws["C6"] = header_info["name"]
    if header_info.get("department"):
        ws["C7"] = header_info["department"]
    if header_info.get("purpose"):
        ws["C8"] = header_info["purpose"]
    if header_info.get("bank_acc"):
        ws["C9"] = header_info["bank_acc"]

    # Periode awal - akhir bulan di E6 & E7
    valid_dates = [d for d in df["date"] if d]
    if valid_dates:
        ref = min(valid_dates)
    else:
        ref = date.today()
    period_start = date(ref.year, ref.month, 1)
    period_end = date(ref.year, ref.month, monthrange(ref.year, ref.month)[1])
    for cell, val in (("E6", period_start), ("E7", period_end)):
        ws[cell].number_format = "d mmm yyyy"
        ws[cell] = val

    n = len(df)
    available_rows = DATA_END_ROW_DEFAULT - DATA_START_ROW + 1

    if n > available_rows:
        extra = n - available_rows
        ws.insert_rows(DATA_END_ROW_DEFAULT + 1, amount=extra)
        for i in range(extra):
            src_row = DATA_END_ROW_DEFAULT
            dst_row = DATA_END_ROW_DEFAULT + 1 + i
            for col in range(2, 6):
                src_cell = ws.cell(row=src_row, column=col)
                dst_cell = ws.cell(row=dst_row, column=col)
                dst_cell.number_format = src_cell.number_format
                dst_cell.font = src_cell.font.copy()
                dst_cell.border = src_cell.border.copy()
                dst_cell.fill = src_cell.fill.copy()
                dst_cell.alignment = src_cell.alignment.copy()
        total_row = TOTAL_ROW_DEFAULT + extra
        data_end_row = DATA_END_ROW_DEFAULT + extra
    else:
        total_row = TOTAL_ROW_DEFAULT
        data_end_row = DATA_END_ROW_DEFAULT

    for i, row in df.iterrows():
        r = DATA_START_ROW + i
        d = row["date"]
        ws.cell(row=r, column=2, value=d)
        if d:
            ws.cell(row=r, column=2).number_format = "d mmm yyyy"
        ws.cell(row=r, column=3, value=row["category"])
        ws.cell(row=r, column=4, value=row["description"])
        ws.cell(row=r, column=5, value=row["nominal"])
        ws.cell(row=r, column=5).number_format = "#,##0"

    ws.cell(row=total_row, column=5, value=f"=SUM(E{DATA_START_ROW}:E{data_end_row})")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
